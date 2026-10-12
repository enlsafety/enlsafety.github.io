#!/usr/bin/env python3
"""Generate non-sensitive fixtures for the real browser parser/OCR integration test.

Requirements:
    Python 3.9+ and: python -m pip install Pillow reportlab openpyxl
    A Korean-capable TrueType (.ttf) font, such as Pretendard Regular or
    NanumGothic. The font must contain Hangul glyphs; a Latin-only font will
    produce invalid OCR fixtures. Do not supply a TTC/OTC collection or CFF OTF.
    No live incident records, credentials, paid services, or external font
    downloads are used by this generator.

Generate fixtures outside the repository:
    python tests/historical-bulk-format-fixtures.py \
      --out-dir /tmp/enl-bulk-format-qa \
      --font /path/to/Pretendard-Regular.ttf

The defaults are ENL_BULK_FIXTURE_DIR or /tmp/enl-bulk-format-qa for output,
then ENL_BULK_QA_FONT or the local font candidates in resolve_font() for the font.

Run the actual Chromium test from the repository root:
    ENL_BULK_FIXTURE_DIR=/tmp/enl-bulk-format-qa \
      CHROMIUM_PATH=/path/to/headless_shell \
      node tests/historical-bulk-formats-live-browser.cjs

The browser test needs Playwright, an installed Chromium binary, and network
access to the configured public PDF.js, SheetJS, Tesseract and OCR model CDNs.
On the managed runtime, ENL_QA_USE_ENV_PROXY=1 uses HTTPS_PROXY and restricts
certificate exceptions to the public-key fingerprint of CODEX_PROXY_CERT (or
SSL_CERT_FILE); it does not disable TLS validation globally.

The JavaScript test creates the XLS binary from rows.json using the actual
SheetJS writer, then reads it through the production extraction path. This
script creates CSV, date-valued XLSX, time-valued XLSX, text PDF, scanned PDF,
PNG, JPEG, WEBP and rows.json. Do not commit these generated binary fixtures.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, time
import json
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FIELDS = [
    '사업장명', '사고일', '사고구분', '사고유형', '사고정도', '사고장소',
    '당시작업', '사고경위', '당시조치', '피해자', '부상내용', 'ERP문서번호',
    'ERP결재일', 'ERP결재상태', '사고번호', '진료비',
]
ROWS = [
    ['테스트CC', '2025-05-10', '대인', '베임', '경미', '3홀', '예초 작업',
     '낫에 왼손 베임', '응급처치', '김00', '왼손 열상', 'ERP-001',
     '2025-05-11', '결재완료', 'OLD-001', 120000],
    ['테스트CC', '2025-05-09', '대인', '베임', '경미', '4홀', '예초 작업',
     '낫에 오른손 베임', '응급처치', '박00', '오른손 열상', 'ERP-002',
     '2025-05-11', '결재완료', 'OLD-002', 90000],
]


def resolve_font(explicit: str | None) -> Path:
    candidates = [explicit] if explicit else [
        os.environ.get('ENL_BULK_QA_FONT'),
        '/root/.local/share/fonts/Pretendard-Regular.ttf',
        '/usr/share/fonts/truetype/nanum/NanumGothic.ttf',
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            path = Path(candidate).resolve()
            if path.suffix.lower() != '.ttf':
                raise ValueError('Supply a Hangul-capable TrueType .ttf font.')
            return path
    raise FileNotFoundError(
        'Korean TTF font not found. Use --font /path/to/Pretendard-Regular.ttf '
        'or set ENL_BULK_QA_FONT.'
    )


def generate(output: Path, font_path: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    with (output / 'two-incidents.csv').open('w', encoding='utf-8-sig', newline='') as file:
        csv.writer(file).writerows([FIELDS] + ROWS)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '사고목록'
    sheet.append(FIELDS)
    for row in ROWS:
        values = list(row)
        for column in (1, 12):
            values[column] = datetime.strptime(values[column], '%Y-%m-%d')
        sheet.append(values)
    for row in sheet.iter_rows(min_row=2):
        for column in (1, 12):
            row[column].number_format = 'm/d/yy'
    workbook.save(output / 'two-incidents.xlsx')

    # A real Excel time serial catches accidental date-only normalization.
    sheet.cell(1, 17, '발생시간')
    for row in (2, 3):
        sheet.cell(row, 17, time(9, 30)).number_format = 'h:mm'
    workbook.save(output / 'time-incidents.xlsx')

    lines = [f'{key}: {value}' for key, value in zip(FIELDS, ROWS[0])]
    pdfmetrics.registerFont(TTFont('PretendardQA', str(font_path)))
    document = canvas.Canvas(str(output / 'text-report.pdf'), pagesize=(700, 980))
    document.setFont('PretendardQA', 17)
    for index, line in enumerate(lines):
        document.drawString(40, 930 - index * 48, line)
    document.save()

    image = Image.new('RGB', (1700, 2200), 'white')
    drawing = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 56)
    for index, line in enumerate(lines):
        drawing.text((70, 70 + index * 125), line, font=font, fill='black')
    image.save(output / 'scan.png')
    image.save(output / 'scan.jpg', quality=95)
    image.save(output / 'scan.webp', quality=95)

    document = canvas.Canvas(str(output / 'scanned-report.pdf'), pagesize=(700, 980))
    document.drawImage(str(output / 'scan.png'), 0, 0, width=700, height=980)
    document.save()
    (output / 'rows.json').write_text(json.dumps([FIELDS] + ROWS, ensure_ascii=False), encoding='utf-8')
    print(f'Created synthetic parser/OCR fixtures in {output.resolve()}')


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--out-dir', type=Path, default=Path(os.environ.get('ENL_BULK_FIXTURE_DIR', '/tmp/enl-bulk-format-qa')))
    parser.add_argument('--font', help='Path to a Korean-capable TrueType font')
    args = parser.parse_args()
    generate(args.out_dir, resolve_font(args.font))


if __name__ == '__main__':
    main()
